"""tracker.csv -> Excel (.xlsx), one row per selected compound.

usage: python build_excel.py tracker.csv out.xlsx

Columns: Compound_number, Compound Name, Patent No., MW, Activity, any extra data
columns (e.g. clogp, clint, papp), Smiles. Workflow columns (batch, reason, flags, ...)
are left out. If the tracker has a "selected" column, only selected rows are kept.
MW is computed from the SMILES when the tracker has none.
"""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Excel column -> tracker columns to take it from (first one that has a value)
MAIN = [('Compound_number', ['compound_number', 'label']),
        ('Compound Name', ['compound_name', 'name', 'iupac_name']),
        ('Patent No.', ['patent_no']),
        ('MW', ['mw']),
        ('Activity', ['activity_display', 'activity', 'activity_raw'])]
SMILES = ['smiles']
WORKFLOW = {'order', 'source_tables', 'pages', 'activity_raw', 'selected', 'reason', 'batch', 'drawn',
            'structure_source', 'confidence', 'validation_status', 'status', 'flags', 'inchikey', 'file',
            'formula', 'notes'}
YES = {'yes', 'y', 'true', '1', 'x'}


def _pick(row, keys):
    for k in keys:
        v = (row.get(k) or '').strip()
        if v:
            return v
    return ''


def _num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() and '.' not in v else f
    except ValueError:
        return v


def _mw(smiles):
    try:
        from rdkit import Chem, RDLogger
        from rdkit.Chem import Descriptors
        RDLogger.DisableLog('rdApp.*')
        m = Chem.MolFromSmiles(smiles)
        return round(Descriptors.MolWt(m), 2) if m else ''
    except ImportError:
        return ''


def build(tracker, out):
    with open(tracker, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
        cols = [c.strip() for c in (rows[0].keys() if rows else [])]
    rows = [{(k or '').strip().lower(): (v or '') for k, v in r.items()} for r in rows]
    if rows and 'selected' in rows[0]:
        rows = [r for r in rows if r['selected'].strip().lower() in YES]
    if rows and 'order' in rows[0]:
        rows.sort(key=lambda r: float(r['order']) if r['order'].strip().replace('.', '', 1).isdigit() else 1e9)
    used = {k for _, ks in MAIN for k in ks} | set(SMILES) | WORKFLOW
    extras = [c for c in cols if c.lower() not in used]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Compounds'
    header = [h for h, _ in MAIN] + extras + ['Smiles']
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        smi = _pick(r, SMILES)
        vals = [_pick(r, ks) for _, ks in MAIN]
        if not vals[3] and smi:
            vals[3] = _mw(smi)
        vals[3] = _num(vals[3]) if isinstance(vals[3], str) else vals[3]
        vals += [_num(r.get(c.lower(), '').strip()) for c in extras]
        ws.append(vals + [smi])
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(c.value or '')) for c in ws[get_column_letter(i)]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = 'A2'
    wb.save(out)
    return len(rows)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    print(build(sys.argv[1], sys.argv[2]), 'compounds ->', sys.argv[2])
